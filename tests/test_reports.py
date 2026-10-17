import json
from datetime import date
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from src.pipeline import (
    EXCEL_SHEET_NAME,
    REPORT_COLUMNS,
    alert_threshold_for_unit,
    export_alerts_json,
    export_excel,
)

REPORT_DATE = date(2026, 10, 16)


@pytest.fixture
def report_df() -> pd.DataFrame:
    return pd.DataFrame(
        {
            "raw_city": ["  dubai$$", "  dubai$$", "san_josé", "~singapore~"],
            "city": ["Dubai", "Dubai", "San José", "Singapore"],
            "latitude": [25.2048, 25.2048, 9.9281, 1.3521],
            "longitude": [55.2708, 55.2708, -84.0907, 103.8198],
            "date": [REPORT_DATE, date(2026, 10, 17), REPORT_DATE, REPORT_DATE],
            "max_temperature": [36.24, 29.5, 27.3, 30.0],
            "total_precipitation": [0.0, 0.0, 12.345, 3.1],
        }
    )[REPORT_COLUMNS]


def test_export_excel_creates_formatted_report(tmp_path: Path, report_df: pd.DataFrame) -> None:
    reports_dir = tmp_path / "reports"

    path = export_excel(report_df, reports_dir, REPORT_DATE)

    assert path == reports_dir / "weather_report_20261016.xlsx"
    sheet = load_workbook(path)[EXCEL_SHEET_NAME]
    assert [cell.value for cell in sheet[1]] == REPORT_COLUMNS
    assert sheet.max_row == len(report_df) + 1
    assert all(cell.font.bold for cell in sheet[1])
    assert all(cell.fill.fill_type == "solid" for cell in sheet[1])
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:G5"
    assert sheet["C2"].number_format == "0.00"
    assert sheet["E2"].number_format == "yyyy-mm-dd"
    assert sheet["F2"].number_format == "0.0"
    assert sheet["G2"].number_format == "0.00"
    assert sheet.column_dimensions["G"].width >= len("total_precipitation")


def test_export_alerts_json_only_includes_rows_above_threshold(
    tmp_path: Path, report_df: pd.DataFrame
) -> None:
    path = export_alerts_json(report_df, 30.0, tmp_path, REPORT_DATE)

    assert path == tmp_path / "weather_alerts_20261016.json"
    alerts = json.loads(path.read_text(encoding="utf-8"))
    assert len(alerts) == 1
    alert = alerts[0]
    assert set(alert) == {"city", "date", "max_temperature", "threshold", "generated_at"}
    assert alert["city"] == "Dubai"
    assert alert["date"] == "2026-10-16"
    assert alert["max_temperature"] == 36.2
    assert alert["threshold"] == 30.0


def test_export_alerts_json_respects_threshold_value(
    tmp_path: Path, report_df: pd.DataFrame
) -> None:
    path = export_alerts_json(report_df, 25.0, tmp_path, REPORT_DATE)

    alerts = json.loads(path.read_text(encoding="utf-8"))
    assert [alert["city"] for alert in alerts] == ["Dubai", "Singapore", "Dubai", "San José"]
    assert all(alert["max_temperature"] > 25.0 for alert in alerts)


def test_export_alerts_json_writes_empty_list_when_nothing_qualifies(
    tmp_path: Path, report_df: pd.DataFrame
) -> None:
    path = export_alerts_json(report_df, 40.0, tmp_path, REPORT_DATE)

    assert json.loads(path.read_text(encoding="utf-8")) == []


def test_alert_threshold_for_unit() -> None:
    assert alert_threshold_for_unit(30.0) == 30.0
    assert alert_threshold_for_unit(30.0, "fahrenheit") == pytest.approx(86.0)
